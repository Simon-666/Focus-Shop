import os
import sys
import json
import re
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EXCEL_FILE = 'products_template.xlsx'
CSV_FILE = 'products_template.csv'
JS_FILE = 'products.js'

def get_categories():
    if os.path.exists(JS_FILE):
        with open(JS_FILE, 'r', encoding='utf-8') as f:
            text = f.read()
        cat_match = re.search(r'const categories = (\[[\s\S]*?\]);', text)
        if cat_match:
            try: return json.loads(cat_match.group(1))
            except: pass
    return [
        {'id': 'all', 'name': 'الكل'},
        {'id': 'photography', 'name': 'التصوير'},
        {'id': 'home', 'name': 'المنزل'},
        {'id': 'camping', 'name': 'التخييم والصيد'},
        {'id': 'clothing', 'name': 'الأحذية والملابس'},
        {'id': 'pc', 'name': 'الكمبيوتر والإنتاجية'}
    ]

def read_from_js():
    if not os.path.exists(JS_FILE):
        return []
    with open(JS_FILE, 'r', encoding='utf-8') as f:
        text = f.read()
    prod_match = re.search(r'const products = (\[[\s\S]*?\]);', text)
    if prod_match:
        try: return json.loads(prod_match.group(1))
        except: pass
    return []

def map_headers(headers):
    mapping = {}
    for idx, h in enumerate(headers):
        hl = h.strip().lower()
        if 'id' in hl and 'category' not in hl and 'قسم' not in hl:
            mapping['id'] = idx
        elif 'title' in hl or 'اسم المنتج' in hl:
            mapping['title'] = idx
        elif 'السعر الأصلي' in hl or 'original' in hl:
            mapping['originalPrice'] = idx
        elif 'السعر الحالي' in hl or hl == 'price' or 'price' in hl:
            if 'original' not in hl and 'أصلي' not in hl:
                mapping['price'] = idx
        elif 'currency' in hl or 'عملة' in hl:
            mapping['currency'] = idx
        elif 'category id' in hl or 'معرف القسم' in hl:
            mapping['category'] = idx
        elif 'category' in hl or 'قسم' in hl:
            if 'name' not in hl and 'اسم' not in hl and 'category' not in mapping:
                mapping['category'] = idx
        elif 'sold' in hl or 'نفدت' in hl or 'مباع' in hl:
            mapping['isSoldOut'] = idx
        elif 'quantity' in hl or 'الكمية' in hl:
            mapping['quantity'] = idx
        elif 'condition' in hl or 'حالة' in hl:
            mapping['condition'] = idx
        elif 'featured' in hl or 'مميز' in hl:
            mapping['featured'] = idx
        elif 'demo' in hl or 'تجريبي' in hl:
            mapping['isDemo'] = idx
        elif 'sort' in hl or 'ترتيب' in hl:
            mapping['sortOrder'] = idx
        elif 'image' in hl or 'صورة' in hl:
            mapping['image'] = idx
        elif 'url' in hl or 'رابط' in hl:
            mapping['url'] = idx
        elif 'description' in hl or 'وصف' in hl:
            mapping['description'] = idx
    return mapping

def read_from_csv():
    if not os.path.exists(CSV_FILE):
        return []
    with open(CSV_FILE, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        rows = [r for r in reader if any(r)]
    if len(rows) < 2:
        return []

    mapping = map_headers(rows[0])
    if 'title' not in mapping or 'price' not in mapping:
        print('Error: Could not identify Title or Price columns in CSV!')
        return []

    products = []
    for r_idx, row in enumerate(rows[1:], 2):
        title = row[mapping['title']].strip() if mapping['title'] < len(row) else ''
        if not title:
            continue

        raw_price = row[mapping['price']] if mapping['price'] < len(row) else '0'
        price = re.sub(r'[^\d]', '', str(raw_price)) or '0'

        orig_price = None
        if 'originalPrice' in mapping and mapping['originalPrice'] < len(row):
            raw_orig = row[mapping['originalPrice']]
            cleaned_orig = re.sub(r'[^\d]', '', str(raw_orig))
            if cleaned_orig and cleaned_orig != '0':
                orig_price = cleaned_orig

        qty = 1
        if 'quantity' in mapping and mapping['quantity'] < len(row):
            try: qty = int(row[mapping['quantity']])
            except: qty = 1

        is_sold = False
        if 'isSoldOut' in mapping and mapping['isSoldOut'] < len(row):
            s_val = str(row[mapping['isSoldOut']]).lower()
            is_sold = 'true' in s_val or 'نعم' in s_val or qty <= 0
        else:
            is_sold = (qty <= 0)

        p_id = 1779378000000 + r_idx
        if 'id' in mapping and mapping['id'] < len(row):
            try: p_id = int(row[mapping['id']])
            except: pass

        category = row[mapping['category']].strip() if 'category' in mapping and mapping['category'] < len(row) else 'photography'
        currency = row[mapping['currency']].strip() if 'currency' in mapping and mapping['currency'] < len(row) else 'د.ع'
        image = row[mapping['image']].strip() if 'image' in mapping and mapping['image'] < len(row) else ''
        url = row[mapping['url']].strip() if 'url' in mapping and mapping['url'] < len(row) else ''
        desc = row[mapping['description']].strip() if 'description' in mapping and mapping['description'] < len(row) else ''
        cond = row[mapping['condition']].strip().lower() if 'condition' in mapping and mapping['condition'] < len(row) else 'new'
        condition = 'used' if 'used' in cond or 'مستخدم' in cond else 'new'

        feat = str(row[mapping['featured']]).lower() if 'featured' in mapping and mapping['featured'] < len(row) else ''
        featured = 'true' in feat or 'نعم' in feat

        dem = str(row[mapping['isDemo']]).lower() if 'isDemo' in mapping and mapping['isDemo'] < len(row) else ''
        is_demo = 'true' in dem or 'نعم' in dem

        sort_order = r_idx - 1
        if 'sortOrder' in mapping and mapping['sortOrder'] < len(row):
            try: sort_order = int(row[mapping['sortOrder']])
            except: pass

        prod = {
            'id': p_id,
            'title': title,
            'price': price,
            'currency': currency or 'د.ع',
            'category': category or 'photography',
            'quantity': qty,
            'isSoldOut': is_sold,
            'image': image,
            'description': desc,
            'url': url,
            'condition': condition,
            'featured': featured,
            'isDemo': is_demo,
            'sortOrder': sort_order
        }
        if orig_price:
            prod['originalPrice'] = orig_price
        products.append(prod)

    return products

def read_from_excel():
    if not os.path.exists(EXCEL_FILE):
        return []
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    ws = wb['المنتجات (Products)'] if 'المنتجات (Products)' in wb.sheetnames else wb.active

    headers = [str(ws.cell(row=1, column=c).value or '').strip() for c in range(1, ws.max_column + 1)]
    mapping = map_headers(headers)
    if 'title' not in mapping or 'price' not in mapping:
        print('Error: Could not identify Title or Price columns in Excel!')
        return []

    products = []
    for r_idx in range(2, ws.max_row + 1):
        title = str(ws.cell(row=r_idx, column=mapping['title'] + 1).value or '').strip()
        if not title:
            continue

        raw_price = ws.cell(row=r_idx, column=mapping['price'] + 1).value
        price = re.sub(r'[^\d]', '', str(raw_price or '0')) or '0'

        orig_price = None
        if 'originalPrice' in mapping:
            raw_orig = ws.cell(row=r_idx, column=mapping['originalPrice'] + 1).value
            if raw_orig:
                cleaned_orig = re.sub(r'[^\d]', '', str(raw_orig))
                if cleaned_orig and cleaned_orig != '0':
                    orig_price = cleaned_orig

        qty = 1
        if 'quantity' in mapping:
            try: qty = int(ws.cell(row=r_idx, column=mapping['quantity'] + 1).value or 1)
            except: qty = 1

        is_sold = False
        if 'isSoldOut' in mapping:
            s_val = str(ws.cell(row=r_idx, column=mapping['isSoldOut'] + 1).value or '').lower()
            is_sold = 'true' in s_val or 'نعم' in s_val or qty <= 0
        else:
            is_sold = (qty <= 0)

        p_id = 1779378000000 + r_idx
        if 'id' in mapping:
            try: p_id = int(ws.cell(row=r_idx, column=mapping['id'] + 1).value or p_id)
            except: pass

        category = str(ws.cell(row=r_idx, column=mapping['category'] + 1).value or 'photography').strip() if 'category' in mapping else 'photography'
        currency = str(ws.cell(row=r_idx, column=mapping['currency'] + 1).value or 'د.ع').strip() if 'currency' in mapping else 'د.ع'
        image = str(ws.cell(row=r_idx, column=mapping['image'] + 1).value or '').strip() if 'image' in mapping else ''
        url = str(ws.cell(row=r_idx, column=mapping['url'] + 1).value or '').strip() if 'url' in mapping else ''
        desc = str(ws.cell(row=r_idx, column=mapping['description'] + 1).value or '').strip() if 'description' in mapping else ''
        cond = str(ws.cell(row=r_idx, column=mapping['condition'] + 1).value or 'new').strip().lower() if 'condition' in mapping else 'new'
        condition = 'used' if 'used' in cond or 'مستخدم' in cond else 'new'

        feat = str(ws.cell(row=r_idx, column=mapping['featured'] + 1).value or '').lower() if 'featured' in mapping else ''
        featured = 'true' in feat or 'نعم' in feat

        dem = str(ws.cell(row=r_idx, column=mapping['isDemo'] + 1).value or '').lower() if 'isDemo' in mapping else ''
        is_demo = 'true' in dem or 'نعم' in dem

        sort_order = r_idx - 1
        if 'sortOrder' in mapping:
            try: sort_order = int(ws.cell(row=r_idx, column=mapping['sortOrder'] + 1).value or sort_order)
            except: pass

        prod = {
            'id': p_id,
            'title': title,
            'price': price,
            'currency': currency or 'د.ع',
            'category': category or 'photography',
            'quantity': qty,
            'isSoldOut': is_sold,
            'image': image,
            'description': desc,
            'url': url,
            'condition': condition,
            'featured': featured,
            'isDemo': is_demo,
            'sortOrder': sort_order
        }
        if orig_price:
            prod['originalPrice'] = orig_price
        products.append(prod)

    return products

def save_to_js(categories, products):
    out = 'const categories = ' + json.dumps(categories, ensure_ascii=False, indent=4) + ';\n\nconst products = ' + json.dumps(products, ensure_ascii=False, indent=4) + ';\n'
    with open(JS_FILE, 'w', encoding='utf-8') as f:
        f.write(out)

def save_to_csv(categories, products):
    cat_map = {c['id']: c['name'] for c in categories}
    headers = [
        'ID (معرف فريد)',
        'اسم المنتج (Title)',
        'السعر الحالي (Price)',
        'السعر الأصلي (Original Price)',
        'العملة (Currency)',
        'معرف القسم (Category ID)',
        'اسم القسم (Category Name)',
        'الكمية المتوفرة (Quantity)',
        'نفدت الكمية؟ (Is Sold Out)',
        'الحالة (Condition: new/used)',
        'مميز في الأعلى؟ (Featured)',
        'تجريبي؟ (Demo)',
        'ترتيب العرض (Sort Order)',
        'رابط الصورة (Image URL)',
        'رابط الشراء الخارجي (Product URL)',
        'وصف المنتج (Description)'
    ]
    rows = [headers]
    for p in products:
        qty = p.get('quantity', 1)
        is_sold = p.get('isSoldOut', False) or (qty <= 0)
        row = [
            str(p.get('id', '')),
            p.get('title', ''),
            str(p.get('price', '')),
            str(p.get('originalPrice', '') or ''),
            p.get('currency', 'د.ع'),
            p.get('category', ''),
            cat_map.get(p.get('category'), p.get('category', '')),
            str(qty),
            'نعم (TRUE)' if is_sold else 'لا (FALSE)',
            p.get('condition', 'new'),
            'نعم (TRUE)' if p.get('featured') else 'لا (FALSE)',
            'نعم (TRUE)' if p.get('isDemo') else 'لا (FALSE)',
            str(p.get('sortOrder', 1)),
            p.get('image', ''),
            p.get('url', ''),
            p.get('description', '')
        ]
        rows.append(row)
    with open(CSV_FILE, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerows(rows)

def save_to_excel(categories, products):
    cat_map = {c['id']: c['name'] for c in categories}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'المنتجات (Products)'
    ws.views.sheetView[0].rightToLeft = True

    headers = [
        ('ID (معرف فريد)', 'id'),
        ('اسم المنتج (Title)', 'title'),
        ('السعر الحالي (Price)', 'price'),
        ('السعر الأصلي (Original Price)', 'originalPrice'),
        ('العملة (Currency)', 'currency'),
        ('معرف القسم (Category ID)', 'category'),
        ('اسم القسم (Category Name)', '_cat_name'),
        ('الكمية المتوفرة (Quantity)', 'quantity'),
        ('نفدت الكمية؟ (Is Sold Out)', 'isSoldOut'),
        ('الحالة (Condition: new/used)', 'condition'),
        ('مميز في الأعلى؟ (Featured)', 'featured'),
        ('تجريبي؟ (Demo)', 'isDemo'),
        ('ترتيب العرض (Sort Order)', 'sortOrder'),
        ('رابط الصورة (Image URL)', 'image'),
        ('رابط الشراء الخارجي (Product URL)', 'url'),
        ('وصف المنتج (Description)', 'description')
    ]

    header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    border_thin = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center', wrap_text=True)

    for col_idx, (header_label, _) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin

    ws.row_dimensions[1].height = 28
    alt_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    sold_fill = PatternFill(start_color='FEF2F2', end_color='FEF2F2', fill_type='solid')

    for r_idx, p in enumerate(products, 2):
        qty = p.get('quantity', 1)
        is_sold = p.get('isSoldOut', False) or (qty <= 0)
        current_fill = sold_fill if is_sold else (alt_fill if r_idx % 2 == 0 else None)

        for c_idx, (_, field) in enumerate(headers, 1):
            if field == '_cat_name': val = cat_map.get(p.get('category'), p.get('category'))
            elif field == 'quantity': val = qty
            elif field == 'isSoldOut': val = 'نعم (TRUE)' if is_sold else 'لا (FALSE)'
            elif field == 'featured' or field == 'isDemo': val = 'نعم (TRUE)' if p.get(field) else 'لا (FALSE)'
            else: val = p.get(field, '') or ''

            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            if current_fill: cell.fill = current_fill
            if field in ['id', 'price', 'originalPrice', 'currency', 'category', 'condition', 'featured', 'isDemo', 'sortOrder', 'quantity', 'isSoldOut']:
                cell.alignment = align_center
            else:
                cell.alignment = align_right

        ws.row_dimensions[r_idx].height = 24

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '').split('\n')[0]
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 45), 14)

    # Guide Sheet
    ws_guide = wb.create_sheet(title='دليل الاستخدام (Guide)')
    ws_guide.views.sheetView[0].rightToLeft = True
    guide_headers = ['اسم العمود (Column)', 'مطلوب؟ (Required)', 'الوصف والخيارات المتاحة (Description & Options)']
    for c_idx, h in enumerate(guide_headers, 1):
        cell = ws_guide.cell(row=1, column=c_idx, value=h)
        cell.fill = PatternFill(start_color='047857', end_color='047857', fill_type='solid')
        cell.font = header_font
        cell.alignment = align_center
        cell.border = border_thin
    ws_guide.row_dimensions[1].height = 26

    guide_data = [
        ('ID (معرف فريد)', 'نعم', 'رقم فريد لكل منتج'),
        ('اسم المنتج (Title)', 'نعم', 'اسم المنتج بالعربية أو الإنجليزية'),
        ('السعر الحالي (Price)', 'نعم', 'السعر كرقم بدون فواصل'),
        ('السعر الأصلي (Original Price)', 'اختياري', 'السعر قبل الخصم'),
        ('العملة (Currency)', 'نعم', 'افتراضياً د.ع'),
        ('معرف القسم (Category ID)', 'نعم', 'photography, home, camping, clothing, pc'),
        ('الكمية المتوفرة (Quantity)', 'نعم', '0 للمباع، 1 لمتبقي قطعة واحدة، 5 لمتوفر'),
        ('نفدت الكمية؟ (Is Sold Out)', 'نعم', 'نعم (TRUE) للخلصان، لا (FALSE) للمتوفر'),
        ('الحالة (Condition)', 'نعم', 'new (جديد) أو used (مستخدم)'),
        ('مميز في الأعلى؟ (Featured)', 'اختياري', 'نعم (TRUE) أو لا (FALSE)'),
        ('تجريبي؟ (Demo)', 'اختياري', 'نعم (TRUE) للعرض فقط'),
        ('ترتيب العرض (Sort Order)', 'اختياري', '1 هو الأول'),
        ('رابط الصورة (Image URL)', 'نعم', 'رابط مباشر للصورة'),
        ('رابط الشراء الخارجي (Product URL)', 'اختياري', 'رابط أمازون أو الموقع الأصلي'),
        ('وصف المنتج (Description)', 'نعم', 'مواصفات وتفاصيل المنتج')
    ]
    for r_idx, row in enumerate(guide_data, 2):
        for c_idx, val in enumerate(row, 1):
            cell = ws_guide.cell(row=r_idx, column=c_idx, value=val)
            cell.border = border_thin
            if c_idx == 2:
                cell.alignment = align_center
                cell.font = Font(name='Calibri', color='DC2626' if val == 'نعم' else '4B5563', bold=(val == 'نعم'))
            else:
                cell.alignment = align_right

    ws_guide.column_dimensions['A'].width = 30
    ws_guide.column_dimensions['B'].width = 16
    ws_guide.column_dimensions['C'].width = 75
    wb.save(EXCEL_FILE)

def main():
    arg = sys.argv[1].lower() if len(sys.argv) > 1 else 'auto'
    categories = get_categories()

    source = None
    if arg in ['csv', '--csv']: source = 'csv'
    elif arg in ['excel', 'xlsx', '--excel']: source = 'excel'
    elif arg in ['js', '--js']: source = 'js'
    else:
        csv_time = os.path.getmtime(CSV_FILE) if os.path.exists(CSV_FILE) else 0
        excel_time = os.path.getmtime(EXCEL_FILE) if os.path.exists(EXCEL_FILE) else 0
        if csv_time > excel_time: source = 'csv'
        else: source = 'excel'

    if source == 'csv':
        print(f'Syncing FROM: {CSV_FILE} ...')
        products = read_from_csv()
        if not products:
            print('Safety Check: No valid products found in CSV! Aborting sync to prevent data loss.')
            return
        valid_prices = [p for p in products if p.get('price') and p['price'] != '0']
        if len(valid_prices) == 0:
            print('Safety Check: All products have price 0 in CSV! Aborting sync to protect your prices.')
            return
        save_to_excel(categories, products)
        save_to_js(categories, products)
        print(f'Done! Successfully synced {len(products)} products from CSV -> Excel and JS.')
    elif source == 'js':
        print(f'Syncing FROM: {JS_FILE} ...')
        products = read_from_js()
        if not products:
            print('Safety Check: No valid products found in JS! Aborting sync to prevent data loss.')
            return
        valid_prices = [p for p in products if p.get('price') and p['price'] != '0']
        if len(valid_prices) == 0:
            print('Safety Check: All products have price 0 in JS! Aborting sync to protect your prices.')
            return
        save_to_excel(categories, products)
        save_to_csv(categories, products)
        print(f'Done! Successfully synced {len(products)} products from JS -> Excel and CSV.')
    else:
        print(f'Syncing FROM: {EXCEL_FILE} ...')
        products = read_from_excel()
        if not products:
            print('Safety Check: No valid products found in Excel! Aborting sync to prevent data loss.')
            return
        valid_prices = [p for p in products if p.get('price') and p['price'] != '0']
        if len(valid_prices) == 0:
            print('Safety Check: All products have price 0 in Excel! Aborting sync to protect your prices.')
            return
        save_to_csv(categories, products)
        save_to_js(categories, products)
        print(f'Done! Successfully synced {len(products)} products from Excel -> CSV and JS.')

if __name__ == '__main__':
    main()
